#!/usr/bin/env python3
"""Load a single Marketshare entry-list file (.csv or .xlsx) into the database.

The workbook importer in scripts/import_marketshare_xlsx/ handles the whole
OneDrive corpus. This is the one-file version, for the sheet somebody has just
filled in and wants in the database today.

    python scripts/load_marketshare_csv.py db/marketshare.db "<file>" \
        --regatta "Regates Royales" --year 2026 --segment Classic \
        --country FRA --dry-run

Drop --dry-run to write. Re-running replaces what it wrote last time rather
than duplicating: everything is tagged
races.source_url = 'marketshare-csv:<regatta> <year>'.

Sailmaker is imported, because on a hand-filled marketshare sheet that column
is the whole point. Where the sheet disagrees with a sailmaker we already hold
for that boat, the sheet is NOT allowed to overwrite it silently - the
disagreement is reported and the existing value kept unless --force-sailmaker.
"""
import argparse, csv, datetime, pathlib, re, shutil, sqlite3, sys, unicodedata

TAG = "marketshare-csv:"

# --- column synonyms, same canon as the workbook importer ----------------------
SYN = {
    "sail_no": ["sail number", "sail no", "sail no.", "sail", "bow number"],
    "boat_name": ["yacht name", "boat name", "boat", "yacht", "team", "vessel"],
    "owner": ["owner's name", "owner name", "owner", "owners name"],
    "skipper": ["skipper", "skipper's name", "skipper name", "sailed by", "sailor", "helm"],
    "boat_type": ["yacht type", "boat type", "type", "design", "model"],
    "division": ["class", "division", "fleet"],
    "sailmaker": ["sailmaker", "sail maker", "sail maker [drop down list]"],
    "sail_material": ["sail material", "material"],
    "sales_lead": ["north sails sales lead", "sales lead"],
    "rep": ["north sails rep onboard", "north sails reps onboard", "rep onboard"],
    "country": ["country code", "country", "nat"],
}

SM_ALIAS = {
    "north": "North Sails", "north sails": "North Sails", "ns": "North Sails",
    "n/s": "North Sails", "3di": "North Sails", "partial": "Partial",
    "doyle": "Doyle", "quantum": "Quantum", "uk": "UK", "uk sails": "UK",
    "one sails": "One", "onesails": "One", "ullman": "Ullman",
    "incidence": "Incidence", "elvstrom": "Elvstrom", "elvstrÃ¶m": "Elvstrom",
    "sanders": "Sanders", "technique voile": "Technique Voile",
    "all purpose": "All Purpose", "evolution": "Evolution", "hyde": "Hyde",
    "banks": "Banks", "unknown": "Unknown", "other": "Other",
}


def clean(s):
    """Strip control bytes, NBSPs and the stray markers hand-kept sheets carry."""
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFC", s)
    s = s.replace(" ", " ").replace("�", "")
    s = re.sub(r"[\x00-\x1f\x7f-\x9f]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def hdr(s):
    return re.sub(r"\s+", " ", clean(s).lower()).strip()


# sail numbers that carry no information
PLACEHOLDER = re.compile(r"^(na|n/a|nan|none|tbc|tbа|\?+|-+|0+|\.+)$", re.I)


def clean_sail(s):
    s = clean(s)
    s = re.sub(r"[^\w /?.-]", " ", s)
    s = re.sub(r"\s+", " ", s).strip(" -.")
    if not s:
        return None
    parts = [p for p in s.split() if not PLACEHOLDER.match(p)]
    s = " ".join(parts).strip()
    return s or None


def nname(s):
    s = clean(s).lower()
    s = re.sub(r"\b(the|a|an)\b", " ", s)
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def nsail(s):
    return re.sub(r"[^A-Z0-9]", "", clean(s).upper()) or None


def sail_core(s):
    s = nsail(s)
    m = re.search(r"(\d{2,6})", s or "")
    return m.group(1) if m else None


def norm_sm(s):
    k = clean(s).lower()
    if not k:
        return None
    if k in SM_ALIAS:
        return SM_ALIAS[k]
    if k.startswith("north"):
        return "North Sails"
    if k.startswith("partial"):
        return "Partial"
    return "Other"


def read_rows(path):
    """-> list of dicts on the canonical field names."""
    path = pathlib.Path(path)
    if path.suffix.lower() in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        # Hand-kept CSVs come out of Excel as cp1252 far more often than UTF-8;
        # decoding one as the other is what turns Sebastien into S<?>bastien.
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        grid = list(csv.reader(text.splitlines()))

    lookup = {}
    for canon, names in SYN.items():
        for n in names:
            lookup.setdefault(hdr(n), canon)

    # find the header row
    hrow, cmap = None, {}
    for i, row in enumerate(grid[:10]):
        m, seen = {}, set()
        for j, c in enumerate(row):
            k = lookup.get(hdr(c))
            if k and k not in seen:
                m[j] = k
                seen.add(k)
        if len(m) > len(cmap):
            hrow, cmap = i, m
    if not cmap or "boat_name" not in cmap.values():
        raise SystemExit("No recognisable header row (need at least a boat-name column)")

    out = []
    for i, row in enumerate(grid[hrow + 1:], hrow + 2):
        rec = {}
        for j, canon in cmap.items():
            v = clean(row[j]) if j < len(row) else ""
            if v:
                rec[canon] = v
        if rec.get("sail_no"):
            rec["sail_no"] = clean_sail(rec["sail_no"])
        if not (rec.get("boat_name") or rec.get("sail_no")):
            continue
        rec["_row"] = i
        out.append(rec)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("db")
    ap.add_argument("file")
    ap.add_argument("--regatta", required=True)
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--segment", default=None,
                    help="Grand Prix | Premier Race | Race | One-Design | Classic | Cruise | Premier Cruise")
    ap.add_argument("--country", default=None)
    ap.add_argument("--region", default=None)
    ap.add_argument("--start-date", default=None)
    ap.add_argument("--end-date", default=None)
    ap.add_argument("--force-sailmaker", action="store_true",
                    help="let the sheet overwrite a sailmaker we already hold")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    rows = read_rows(a.file)
    print(f"read {len(rows)} entry rows from {pathlib.Path(a.file).name}")

    dbp = pathlib.Path(a.db)
    if not a.dry_run:
        bak = dbp.with_suffix(f".db.bak-csv-{datetime.datetime.now():%H%M%S}")
        shutil.copy2(dbp, bak)
        print("backup ->", bak.name)
    con = sqlite3.connect(dbp)
    cur = con.cursor()
    tag = f"{TAG}{a.regatta} {a.year}"

    # clear a previous run of this same event
    old = [r[0] for r in cur.execute("select id from races where source_url=?", (tag,))]
    if old:
        q = ",".join("?" * len(old))
        cur.execute(f"delete from race_entries where race_id in ({q})", old)
        cur.execute(f"delete from races where id in ({q})", old)
        print(f"cleared {len(old)} race(s) from a previous load of this event")

    # --- boat index
    by_sail, by_core, by_name = {}, {}, {}
    def index(bid, sn, bn):
        k = nsail(sn)
        if k:
            by_sail.setdefault(k, bid)
        c = sail_core(sn)
        if c:
            by_core.setdefault(c, []).append((bid, nname(bn)))
        if bn:
            by_name.setdefault(nname(bn), []).append((bid, nsail(sn)))
    for bid, sn, bn in cur.execute("select id, sail_no, boat_name from boats"):
        index(bid, sn, bn)

    def match(r):
        s = nsail(r.get("sail_no"))
        if s and s in by_sail:
            return by_sail[s]
        nm = nname(r.get("boat_name"))
        c = sail_core(r.get("sail_no"))
        if c and c in by_core:
            for bid, bn in by_core[c]:
                if bn and bn == nm:
                    return bid
        if nm and nm in by_name:
            cands = by_name[nm]
            if len(cands) == 1:
                return cands[0][0]
            for bid, sn2 in cands:
                if c and sn2 and c in sn2:
                    return bid
        return None

    sm_id = {n.lower(): i for i, n in cur.execute("select id,name from sailmakers")}
    def sailmaker_id(v):
        if not v:
            return None
        if v.lower() not in sm_id:
            cur.execute("insert into sailmakers(name,notes) values(?,?)",
                        (v, "added by marketshare CSV load"))
            sm_id[v.lower()] = cur.lastrowid
        return sm_id[v.lower()]

    # --- regatta / event / race
    rid = cur.execute("select id from regattas where lower(name)=lower(?)",
                      (a.regatta,)).fetchone()
    if rid:
        rid = rid[0]
        if a.segment:
            cur.execute("update regattas set segment=? where id=?", (a.segment, rid))
    else:
        cur.execute("insert into regattas(name,category,region,segment,country)"
                    " values(?,?,?,?,?)",
                    (a.regatta, "Marketshare", a.region or "Non-UK", a.segment, a.country))
        rid = cur.lastrowid
        print(f"created regatta {a.regatta!r}"
              + (f" in segment {a.segment}" if a.segment else ""))

    ev = cur.execute("select id from events where regatta_id=? and season_year=?",
                     (rid, a.year)).fetchone()
    if ev:
        eid = ev[0]
    else:
        cur.execute("insert into events(regatta_id,season_year,start_date,end_date,"
                    "source_url,notes) values(?,?,?,?,?,?)",
                    (rid, a.year, a.start_date, a.end_date, tag,
                     f"entry list loaded from {pathlib.Path(a.file).name}"))
        eid = cur.lastrowid
    cur.execute("insert into races(event_id,race_name,race_number,race_date,status,"
                "source_url) values(?,?,?,?,?,?)",
                (eid, "Entry List", 1, a.start_date, "entry-list", tag))
    race_id = cur.lastrowid

    stat = {"matched": 0, "new boats": 0, "duplicate rows": 0,
            "sailmaker recorded": 0, "sailmaker blank": 0, "sailmaker disagreed": 0}
    seen, disagreements = set(), []
    for r in rows:
        bid = match(r)
        if bid:
            stat["matched"] += 1
        else:
            cur.execute("insert into boats(sail_no,boat_name,boat_type) values(?,?,?)",
                        (r.get("sail_no"), r.get("boat_name"), r.get("boat_type")))
            bid = cur.lastrowid
            index(bid, r.get("sail_no"), r.get("boat_name"))
            stat["new boats"] += 1
        if bid in seen:
            stat["duplicate rows"] += 1
            continue
        seen.add(bid)

        sm = norm_sm(r.get("sailmaker"))
        if sm:
            stat["sailmaker recorded"] += 1
            prev = cur.execute(
                "select s.name from race_entries re join sailmakers s on s.id=re.sailmaker_id"
                " where re.boat_id=? and re.sailmaker_id is not null"
                " order by re.id desc limit 1", (bid,)).fetchone()
            if prev and prev[0] != sm:
                stat["sailmaker disagreed"] += 1
                disagreements.append((r.get("boat_name"), prev[0], sm))
                if not a.force_sailmaker:
                    sm = prev[0]
        else:
            stat["sailmaker blank"] += 1

        cur.execute(
            "insert into race_entries(race_id,boat_id,class,sail_no_used,boat_name_used,"
            "boat_type_used,owner_name_used,skipper_name_used,sailmaker_id,source)"
            " values(?,?,?,?,?,?,?,?,?,?)",
            (race_id, bid, r.get("division"), r.get("sail_no"), r.get("boat_name"),
             r.get("boat_type"), r.get("owner"), r.get("skipper"),
             sailmaker_id(sm), tag))

    for k, v in stat.items():
        print(f"  {k:<20} {v}")
    if disagreements:
        print(f"\n  sheet disagreed with a sailmaker we hold "
              f"({'applied' if a.force_sailmaker else 'KEPT OURS'}):")
        for bn, was, now in disagreements[:20]:
            print(f"    {bn}: held {was}, sheet says {now}")

    if stat["sailmaker recorded"] == 0:
        print("\n  NOTE: not one row carries a sailmaker, so this file adds fleet "
              "size but no market share.")
    elif stat["sailmaker blank"] > stat["sailmaker recorded"] * 3:
        pct = 100 * stat["sailmaker recorded"] / max(len(seen), 1)
        print(f"\n  NOTE: sailmaker is filled in on {pct:.0f}% of the fleet, so any "
              f"share computed off this event rests on a thin sample.")

    if a.dry_run:
        con.rollback()
        print("\nDRY RUN - nothing written")
    else:
        con.commit()
        print("\ncommitted. Rebuild the dashboard:")
        print("  python scripts/export_dashboard_data.py db/marketshare.db dashboard/data.json")
        print("  python scripts/build_dashboard.py")


if __name__ == "__main__":
    main()
