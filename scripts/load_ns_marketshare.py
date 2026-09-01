#!/usr/bin/env python3
"""
Load the North Sails market-share working sheet.

This is the only source in the project that states, first-hand, WHAT SAILS A
BOAT IS USING. Everything else records who raced where; this records the thing
market share is actually made of.

Be clear about what it is, though: a partly-filled working sheet combined from
several people's market-share exercises, not a CRM export. Of 862 boats it
names a sailmaker for 121. Sail Material, NTG Masts and Rep Onboard are empty
columns. It roughly doubles sailmaker coverage - it does not solve it.

THREE THINGS IN THE SHEET THAT A NAIVE READ GETS WRONG:

1. A MAKER'S NAME IN THE "Partial Inventory" COLUMN IS A SECOND MAKER, not a
   yes/no flag. GLADIATOR reads Sailmaker=Quantum, Partial=North: the boat's
   main inventory is Quantum and it carries some North. Those split-inventory
   boats are the most useful rows in the file - a boat already flying some of
   your sails is the shortest conquest there is - so both makers are recorded,
   the second marked partial.

2. Sailmaker="Partial" (3 boats) names no maker at all. It means a partial
   inventory of something unrecorded, so it is stored as a flag, never as a
   sailmaker.

3. "Doyle/north" is one boat with two makers, and "patial" is a typo for
   partial. Both are handled explicitly rather than being allowed through as
   invented sailmaker names.

MATCHING. Sail numbers are compared with punctuation and spacing removed
("GBR 8414R" and "GBR8414R" are one boat). Where that fails, a boat name is
used only when it is unambiguous on BOTH sides - one boat in the database, one
row in the sheet - because attaching an inventory to the wrong boat is worse
than not attaching it.

Boats the sheet knows and the database does not are NOT created by default:
they have no race record, so inventing them would inflate the fleet the market
share is measured against. --create-missing adds them if that is what you want.

Usage:
  python3 load_ns_marketshare.py <db.sqlite> <sheet.xlsx> [--dry-run]
                                 [--create-missing]
"""
import re
import sys
import argparse
import sqlite3
from collections import defaultdict

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from build_db import (get_or_create_sailmaker, get_or_create_owner,
                      get_or_create_boat, norm, norm_upper)

import openpyxl

SOURCE = "ns:marketshare-sheet"

# Values that are a flag, not a maker.
PARTIAL_WORDS = {"partial", "patial", "part", "partial inventory"}


def key(s):
    """Sail numbers differ only in spacing and punctuation between sources."""
    return re.sub(r"[^A-Z0-9]", "", (str(s or "")).upper())


def cell(row, name):
    v = row.get(name)
    v = None if v is None else str(v).strip()
    return v or ""


def makers_from(sailmaker_cell, partial_cell):
    """-> [(maker, is_partial)], plus a flag when a partial is named but its
    maker is not. Never invents a maker out of a flag word."""
    out, partial_unknown = [], False

    sm = sailmaker_cell.strip()
    if sm.lower() in PARTIAL_WORDS:
        partial_unknown = True          # "Partial" alone names nobody
    elif sm:
        for part in re.split(r"\s*[/,&]\s*", sm):   # "Doyle/north" is two makers
            if part.strip():
                out.append((part.strip(), False))

    p = partial_cell.strip()
    if p:
        if p.lower() in PARTIAL_WORDS:
            # a yes/no flag: the maker already named is the partial one
            out = [(m, True) for m, _ in out] or out
            if not out:
                partial_unknown = True
        else:
            # a MAKER's name here is a second, partial inventory
            for part in re.split(r"\s*[/,&]\s*", p):
                if part.strip():
                    out.append((part.strip(), True))
    return out, partial_unknown


def main():
    p = argparse.ArgumentParser()
    p.add_argument("db")
    p.add_argument("xlsx")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--create-missing", action="store_true")
    args = p.parse_args()

    ws = openpyxl.load_workbook(args.xlsx, data_only=True).worksheets[0]
    hdr = [str(c.value or "").strip() for c in ws[1]]
    rows = [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)
            if any(c.value is not None for c in r)]

    conn = sqlite3.connect(args.db)
    cur = conn.cursor()
    # The sheet lists some boats twice, and INSERT OR IGNORE cannot dedupe
    # without something to violate - so give it one. Earlier loads left exact
    # duplicates behind, which have to go before the index will build.
    dup = cur.execute(
        "DELETE FROM boat_sailmaker_history WHERE id NOT IN ("
        "  SELECT MIN(id) FROM boat_sailmaker_history"
        "  GROUP BY boat_id, sailmaker_id, IFNULL(source,''), IFNULL(confidence,''))"
    ).rowcount
    if dup:
        print(f"removed {dup} pre-existing duplicate sailmaker row(s)")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_bsh_unique "
                "ON boat_sailmaker_history(boat_id, sailmaker_id, source, confidence)")
    cur.execute("""CREATE TABLE IF NOT EXISTS boat_crm (
        boat_id INTEGER PRIMARY KEY, lead_rep TEXT, contacted_by TEXT,
        in_cs INTEGER, tag TEXT, notes TEXT, last_updated TEXT,
        boat_captain TEXT, programme_manager TEXT)""")

    by_sail = defaultdict(list)
    by_name = defaultdict(list)
    for bid, sail, nm in cur.execute(
            "SELECT id, sail_no, boat_name FROM boats"):
        if sail:
            by_sail[key(sail)].append(bid)
        if nm:
            by_name[nm.strip().upper()].append(bid)

    sheet_name_counts = defaultdict(int)
    for r in rows:
        sheet_name_counts[cell(r, "Yacht Name").upper()] += 1

    n_sail = n_name = n_unmatched = 0
    n_hist = n_partial = n_lead = n_created = 0
    unmatched_with_maker, unknown_partial = [], []

    for r in rows:
        sail_raw = cell(r, "Sail Number")
        name_raw = cell(r, "Yacht Name")
        k = key(sail_raw)
        boat_id = None

        if k and len(by_sail.get(k, [])) == 1:
            boat_id = by_sail[k][0]
            n_sail += 1
        else:
            nm = name_raw.upper()
            # a name match only counts when it is unique on BOTH sides
            if nm and len(by_name.get(nm, [])) == 1 and sheet_name_counts[nm] == 1:
                boat_id = by_name[nm][0]
                n_name += 1

        makers, partial_unknown = makers_from(cell(r, "Sailmaker"),
                                              cell(r, "Partial Inventory"))

        if boat_id is None:
            n_unmatched += 1
            if makers:
                unmatched_with_maker.append((sail_raw, name_raw,
                                             ", ".join(m for m, _ in makers)))
            if args.create_missing and not args.dry_run:
                boat_id = get_or_create_boat(cur, sail_raw, name_raw,
                                             cell(r, "Yacht Type") or None)
                owner = get_or_create_owner(cur, cell(r, "Owner Name"))
                if boat_id and owner:
                    cur.execute("UPDATE boats SET current_owner_id = ? WHERE id = ?",
                                (owner, boat_id))
                n_created += 1
            if boat_id is None:
                continue

        for maker, is_partial in makers:
            sm_id = get_or_create_sailmaker(cur, maker)
            if not sm_id:
                continue
            if not args.dry_run:
                cur.execute(
                    "INSERT OR IGNORE INTO boat_sailmaker_history "
                    "(boat_id, sailmaker_id, source, confidence) VALUES (?,?,?,?)",
                    (boat_id, sm_id, SOURCE, "partial" if is_partial else "stated"))
            n_hist += 1
            if is_partial:
                n_partial += 1

        if partial_unknown:
            unknown_partial.append((sail_raw, name_raw))

        lead = cell(r, "North Sails Sales Lead")
        if lead and not args.dry_run:
            cur.execute("INSERT INTO boat_crm (boat_id, lead_rep) VALUES (?,?) "
                        "ON CONFLICT(boat_id) DO UPDATE SET lead_rep = excluded.lead_rep",
                        (boat_id, lead))
        if lead:
            n_lead += 1

    print(f"{len(rows)} sheet row(s)")
    print(f"  matched on sail number : {n_sail}")
    print(f"  matched on unique name : {n_name}")
    print(f"  unmatched              : {n_unmatched}"
          + (f"  ({n_created} created)" if n_created else ""))
    print(f"\n{n_hist} sailmaker record(s) written, of which {n_partial} partial.")
    print(f"{n_lead} sales lead(s) written.")
    if unmatched_with_maker:
        print(f"\n{len(unmatched_with_maker)} row(s) name a sailmaker but match no boat "
              f"- these are the ones worth chasing:")
        for s, n, m in unmatched_with_maker[:12]:
            print(f"    {s:<12} {n[:26]:<26} {m}")
    if unknown_partial:
        print(f"\n{len(unknown_partial)} boat(s) marked partial with no maker named "
              f"- recorded as a flag only, not as a sailmaker:")
        for s, n in unknown_partial[:6]:
            print(f"    {s:<12} {n}")

    if args.dry_run:
        conn.rollback()
        print("\n(dry run - nothing written)")
    else:
        conn.commit()
        print("\ncommitted")
    conn.close()


if __name__ == "__main__":
    main()
