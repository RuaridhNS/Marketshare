#!/usr/bin/env python3
"""
Load RORC Season's Points Championship tables (2023-2025).

RORC results moved to sailracehq.com, whose robots.txt disallows this project's
crawler, so these years were a hole. These workbooks come from the user's own
browser session instead.

WHAT THE DATA IS. Each sheet is one (season, class) standings table. The race
columns hold POINTS, not finishing positions, and the vocabulary is small:

    "0 DNC"     did not compete      -> NOT an entry, skipped
    "82.8"      raced                -> an entry
    "95.2+10"   raced, plus a bonus  -> an entry
    "10 RET"    started, retired     -> an entry, status retired
    "0 DNF"     started, no finish   -> an entry, status retired

That DNC rule is the whole point: it turns a standings table into a per-race
entry list, which is what market share is counted from.

THREE DEFECTS IN THE SOURCE, all reported rather than papered over:

1. TypeOwner is two fields in one cell, with no delimiter and sometimes a TCC
   in the middle - "SUN FAST 3200 R2 1.90 proto The Goodhew family". Boats
   already known by sail number do not need it parsed at all, which covers
   about two thirds. For the rest the longest known boat type is matched off
   the front and the remainder taken as the owner; anything that fails to
   match is stored whole and flagged, never guessed at.
2. Some sail numbers arrive without a country prefix ("3", "49"), which cannot
   identify a boat. Those rows are skipped and counted.
3. Several sheets are empty, marked NOT RETRIEVED by the extraction. Listed on
   every run so the gap stays visible.

Non-IRC classes (MOCRA, Class 40) are skipped: this tool is IRC-only.

Usage:
  python3 load_rorc_season_points.py <db.sqlite> <workbook.xlsx> [...] [--dry-run]
"""
import re
import sys
import argparse
import sqlite3
import unicodedata
from collections import defaultdict

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from build_db import (get_or_create_owner, get_or_create_boat, get_or_create_regatta,
                      get_or_create_event, create_race, norm, norm_upper)

import openpyxl

SHEET = re.compile(r"^(\d{4})\s+(.*)$")
# A sail number is a country prefix and digits. Matching "any letters then
# digits" is not good enough here: boat TYPES take that shape too (JPK1180,
# X-50, NMYD 54), and TypeOwner is the neighbouring column, so a loose pattern
# mistakes every one of them for a misplaced sail number.
COUNTRY = (r"GBR|IRL|FRA|GER|NED|BEL|DEN|SWE|NOR|FIN|EST|LAT|LTU|POL|CZE|AUT|"
           r"SUI|ITA|ESP|POR|GRE|TUR|CRO|SLO|HUN|RUS|USA|CAN|MEX|BRA|ARG|CHI|"
           r"AUS|NZL|JPN|CHN|HKG|SIN|UAE|RSA|MON|MHL|CAY|ANT|BVI|ISV|NCA|SAF|"
           r"SVK|UKR|MLT|CYP|ISR|ISL|LUX|ROU|BUL|SRB|MNE|BLR|MDA|GEO|KAZ|"
           r"URU|PER|COL|VEN|CRC|PAN|DOM|PUR|JAM|TRI|BAR|BAH|BER|KOR|THA|IND")
SAIL_OK = re.compile(rf"^(?:{COUNTRY})\s?-?\d{{1,6}}[A-Z]?$", re.I)
SAIL_LEAD = re.compile(rf"^((?:{COUNTRY})\s?-?\d{{1,6}}[A-Z]?)\s+(.*)$", re.I)
DNC = re.compile(r"^\s*0\s*DNC\s*$", re.I)
RETIRED = re.compile(r"\b(RET|DNF|DNS|RTD)\b", re.I)
DISQ = re.compile(r"\b(DSQ|BFD|UFD|OCS|SCP|NSC)\b", re.I)
NON_IRC = re.compile(r"\b(mocra|class\s*40|multihull)\b", re.I)
NOT_RETRIEVED = re.compile(r"not\s+retrieved", re.I)
# a trailing "(IRC 2H Nationals Race 1)" is scoring context, not the race name
PAREN_TAIL = re.compile(r"\s*\([^)]*\)\s*$")
TCC_IN = re.compile(r"\b[0-2]\.\d{2}\b")

MOJIBAKE = re.compile(r"[Â-Ãâ][\x80-\xbf]")


def demojibake(s):
    if not s or not MOJIBAKE.search(s):
        return s
    try:
        return s.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s


def clean(s):
    return re.sub(r"\s+", " ", demojibake(str(s or "")).replace("�", "")).strip()


def split_type_owner(raw, types_by_len):
    """Peel the longest known boat type off the front; the rest is the owner.

    Returns (boat_type, owner, confident). Never invents a split - an
    unrecognised string comes back whole with confident=False so the caller can
    store it verbatim and flag it.
    """
    s = clean(raw)
    if not s:
        return None, None, True
    flat = re.sub(r"[^a-z0-9]", "", s.lower())
    for t, tflat in types_by_len:
        if tflat and flat.startswith(tflat):
            rest = s[len(s) - (len(s) - _prefix_chars(s, len(tflat))):] if False else None
            # walk the raw string until we have consumed tflat's worth of
            # alphanumerics, so spacing and punctuation differences do not
            # throw the offset out
            seen = 0
            cut = 0
            for i, ch in enumerate(s):
                if ch.isalnum():
                    seen += 1
                cut = i + 1
                if seen >= len(tflat):
                    break
            rest = s[cut:].strip()
            rest = TCC_IN.sub("", rest).strip()      # drop a stray rating
            rest = re.sub(r"^(proto|custom|\(\d+\))\s*", "", rest, flags=re.I).strip()
            return t, (rest or None), True
    return s, None, False


def _prefix_chars(s, n):
    return n


def status_of(cell):
    c = clean(cell)
    if DISQ.search(c):
        return "disqualified"
    if RETIRED.search(c):
        return "retired"
    return "finished"


def points_of(cell):
    m = re.match(r"^\s*([\d.]+)", clean(cell))
    if not m:
        return None
    try:
        base = float(m.group(1))
    except ValueError:
        return None
    bonus = re.search(r"\+\s*([\d.]+)", clean(cell))
    return base + float(bonus.group(1)) if bonus else base


def main():
    p = argparse.ArgumentParser()
    p.add_argument("db")
    p.add_argument("xlsx", nargs="+")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()

    conn = sqlite3.connect(args.db)
    cur = conn.cursor()

    types_by_len = sorted(
        ((t, re.sub(r"[^a-z0-9]", "", t.lower()))
         for (t,) in cur.execute(
             "SELECT DISTINCT boat_type FROM boats WHERE boat_type IS NOT NULL")),
        key=lambda x: -len(x[1]))
    known_sails = {s for (s,) in cur.execute(
        "SELECT sail_no FROM boats WHERE sail_no IS NOT NULL")}

    n_entries = 0
    per_class = defaultdict(int)
    empty_sheets, skipped_non_irc = [], []
    bad_sail, unparsed_type, new_boats = [], [], set()
    realigned = 0

    for path in args.xlsx:
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            title = ws.title.strip()
            if title.lower() in ("readme", "race key"):
                continue
            m = SHEET.match(title)
            # the 2025 workbook has one sheet named for the class alone, with
            # the season in a banner row above the header
            if m:
                year, cls = int(m.group(1)), m.group(2).strip()
            else:
                banner = " ".join(clean(c.value) for c in ws[1])
                ym = re.search(r"(20\d{2})", banner)
                if not ym:
                    empty_sheets.append(f"{title} (no season year)")
                    continue
                year, cls = int(ym.group(1)), title
            if NON_IRC.search(cls):
                skipped_non_irc.append(f"{year} {cls}")
                continue

            # find the header row: the one starting with Rank
            hdr_row = None
            for r in range(1, min(ws.max_row, 8) + 1):
                if clean(ws.cell(r, 1).value).lower() == "rank":
                    hdr_row = r
                    break
            if hdr_row is None:
                note = " ".join(clean(c.value) for row in ws.iter_rows(max_row=3)
                                for c in row)
                empty_sheets.append(f"{year} {cls}" +
                                    (" (NOT RETRIEVED)" if NOT_RETRIEVED.search(note) else ""))
                continue

            hdr = [clean(ws.cell(hdr_row, c).value) for c in range(1, ws.max_column + 1)]
            # The 2025 workbook has a different shape: two header rows, each
            # race spanning a Score/Code pair, and Type and Owner already in
            # their own columns. Detect it rather than assume one layout.
            sub = ([clean(ws.cell(hdr_row + 1, c).value)
                    for c in range(1, ws.max_column + 1)]
                   if hdr_row < ws.max_row else [])
            paired = any(x.lower() == "code" for x in sub)
            if paired:
                for i in range(1, len(hdr)):       # race name spans its pair
                    if not hdr[i] and sub[i].lower() in ("score", "code"):
                        hdr[i] = hdr[i - 1]
            try:
                so = hdr.index("SailNo") if "SailNo" in hdr else hdr.index("Sail No")
            except ValueError:
                empty_sheets.append(f"{year} {cls} (no SailNo column)")
                continue
            to = next((i for i, h in enumerate(hdr) if h in ("TypeOwner", "Type")), None)
            oc = next((i for i, h in enumerate(hdr) if h == "Owner"), None)
            data_start = hdr_row + (2 if paired else 1)
            skip_hdr = ("bonus", "total", "owner", "check", "rank", "boat",
                        "sail no", "sailno", "type", "typeowner", "")
            if paired:
                race_cols, i = [], 0
                while i < len(hdr):
                    if (hdr[i] and hdr[i].lower() not in skip_hdr
                            and sub[i].lower() == "score" and i + 1 < len(hdr)
                            and sub[i + 1].lower() == "code"):
                        race_cols.append((i, hdr[i], i + 1))
                        i += 2
                    else:
                        i += 1
            else:
                race_cols = [(i, h, None) for i, h in enumerate(hdr)
                             if i > max(so, to or so) and h
                             and h.lower() not in skip_hdr]
            is_2h = "two-handed" in cls.lower() or "2h" in cls.lower()

            for row in ws.iter_rows(min_row=data_start, values_only=True):
                if row[0] is None:
                    continue
                boat_name = clean(row[1]) if len(row) > 1 else ""
                sail = norm_upper(clean(row[so])) or ""
                raw_to = clean(row[to]) if to is not None and to < len(row) else ""
                # Some rows are shifted a column: the sail cell holds a scrap of
                # the boat name and the real sail number leads TypeOwner.
                if not SAIL_OK.match(sail):
                    lead = SAIL_LEAD.match(raw_to)
                    if lead:
                        if sail:
                            boat_name = f"{boat_name} {sail}".strip()
                        sail = norm_upper(lead.group(1))
                        raw_to = lead.group(2).strip()
                        realigned += 1
                    else:
                        bad_sail.append(f"{year} {cls}: {boat_name[:26]!r} sail={sail!r}")
                        continue

                btype = owner = None
                if oc is not None:              # Type and Owner already split
                    btype = TCC_IN.sub("", raw_to).strip() or None
                    owner = clean(row[oc]) if oc < len(row) else None
                    if sail not in known_sails:
                        new_boats.add(sail)
                elif raw_to and sail not in known_sails:
                    btype, owner, ok = split_type_owner(raw_to, types_by_len)
                    if not ok:
                        unparsed_type.append(f"{sail} {btype[:44]!r}")
                    new_boats.add(sail)

                boat_id = get_or_create_boat(cur, sail, boat_name, btype)
                if boat_id is None:
                    continue
                owner_id = get_or_create_owner(cur, owner) if owner else None

                for ci, raw_name, code_i in race_cols:
                    cell = clean(row[ci]) if ci < len(row) else ""
                    code = (clean(row[code_i])
                            if code_i is not None and code_i < len(row) else "")
                    if code:
                        if re.fullmatch(r"DNC", code, re.I):
                            continue        # did not compete: not an entry
                        cell = (cell + " " + code).strip()
                    if not cell or DNC.match(cell):
                        continue
                    # 2025 prefixes every race "R15: ..."; strip it, or the same
                    # race lands under a different regatta from the 2023-24 file
                    race_name = re.sub(r"^R\d+\s*:\s*", "", raw_name).strip()
                    race_name = PAREN_TAIL.sub("", race_name).strip()
                    regatta = race_name if race_name.lower().startswith(("rorc", "rolex"))\
                        else f"RORC {race_name}"
                    regatta_id = get_or_create_regatta(cur, regatta, "RORC")
                    event_id = get_or_create_event(cur, regatta_id, year)
                    race_id = create_race(cur, event_id, race_name, status="confirmed",
                                          class_label=cls)
                    cur.execute(
                        "INSERT OR REPLACE INTO race_entries "
                        "(race_id, boat_id, class, sail_no_used, boat_name_used, "
                        " boat_type_used, owner_id, owner_name_used, status, points, "
                        " comments, tag, source) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (race_id, boat_id, cls, sail, boat_name, btype, owner_id, owner,
                         status_of(cell), points_of(cell), cell,
                         "2H" if is_2h else None, "rorc:season-points"))
                    n_entries += 1
                    per_class[(year, cls)] += 1

    print(f"{n_entries} entr(ies) loaded\n")
    for (year, cls), n in sorted(per_class.items()):
        print(f"  {year}  {cls:<22} {n:>5}")

    print(f"\n{len(new_boats)} boat(s) new to the database.")
    if empty_sheets:
        print(f"\n{len(empty_sheets)} sheet(s) held NO DATA - still missing from RORC:")
        for s in empty_sheets:
            print(f"    {s}")
    if skipped_non_irc:
        print(f"\nskipped as non-IRC: {', '.join(skipped_non_irc)}")
    if bad_sail:
        print(f"\n{len(bad_sail)} row(s) skipped - sail number cannot identify a boat:")
        for s in bad_sail[:6]:
            print(f"    {s}")
        if len(bad_sail) > 6:
            print(f"    ... and {len(bad_sail)-6} more")
    if unparsed_type:
        print(f"\n{len(unparsed_type)} boat type(s) not recognised - stored whole, "
              f"owner NOT guessed:")
        for s in unparsed_type[:6]:
            print(f"    {s}")
        if len(unparsed_type) > 6:
            print(f"    ... and {len(unparsed_type)-6} more")

    if args.dry_run:
        conn.rollback()
        print("\n(dry run - nothing written)")
    else:
        conn.commit()
        print("\ncommitted")
    conn.close()


if __name__ == "__main__":
    main()
